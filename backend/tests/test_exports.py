"""Tests for safe exports (CSV/XLSX) and formula injection neutralization."""

import pytest
import io
import openpyxl
import pandas as pd
from app.services.export_service import export_to_csv, export_to_xlsx

def test_export_csv_formula_neutralization():
    df = pd.DataFrame({
        "User": ["Alice", "Bob", "Charlie"],
        "Formula_Col": ["=SUM(A1:A10)", "@Dangerous", "+100-50"],
        "Normal_Col": ["Safe text", "123", ""]
    })
    csv_bytes = export_to_csv(df)
    csv_text = csv_bytes.decode("utf-8-sig")

    # Formula injection characters MUST be neutralized with single apostrophe
    assert "'=SUM(A1:A10)" in csv_text
    assert "'@Dangerous" in csv_text
    assert "'+100-50" in csv_text
    assert "Safe text" in csv_text

def test_export_xlsx_roundtrip():
    df = pd.DataFrame({
        "Branch": ["North", "South"],
        "Revenue": [15000.0, 24000.0],
        "Payload": ["=1+1", "Normal"]
    })
    xlsx_bytes = export_to_xlsx(df, sheet_title="TestExport")

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["TestExport"]
    assert ws.cell(row=1, column=1).value == "Branch"
    assert ws.cell(row=2, column=2).value == 15000.0
    # Payload cell starting with = should be neutralized
    assert ws.cell(row=2, column=3).value == "'=1+1"
